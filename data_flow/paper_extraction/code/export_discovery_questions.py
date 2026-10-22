#!/usr/bin/env python3
"""Exports one row per question against the five discovery axes, as exports/discovery/QUESTIONS.xlsx.

    python3 scripts/export_discovery_questions.py

This is the per-question companion to CATEGORIES.xlsx. That workbook has one tab per axis listing
the CATEGORIES; this one has one row per QUESTION carrying its value on each axis plus the reason
it was placed there:

    question_id | subject | year | question_text | options | official_answer
      | E1 question type  | E1 why
      | E2 format         | E2 why
      | E3 content shape  | E3 why
      | E4 thinking       | E4 why
      | E5 trap           | E5 why

Where the values come from, so each column can be checked:

  * the axis VALUES are inverted out of `exports/discovery/final/<tag>.json` -- each category there
    carries a `question_ids` list, so the per-question assignment already existed and is read, not
    recomputed. Two axes are legitimately multi-valued and are joined with ";": E1, because 30
    questions are built two ways at once (a plain MCQ that is also a figure to read), and E5,
    because that tag runs once per WRONG OPTION, so one question can carry several traps.
  * the WHY text is taken from the same files' `examples` entries, which pair a question_id with a
    written reason. Only 240 of the 2,900 (580 questions x 5 axes) have one, because examples were
    written to illustrate a category rather than to cover every question. Missing cells are left
    EMPTY and the count is printed, rather than filled with a restatement of the category
    definition -- a generic sentence in a per-question column would look like evidence while
    carrying none.
  * question text, options, subject, year and the official answer come from `papers/records/`.

Nothing here is a judgement: every cell is copied from an existing file or left blank.
"""

import argparse
import collections
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import runner_core as rc  # noqa: E402
import question_assets as qa  # noqa: E402

AXES = [
    ("E1", "question_type", "How is the question BUILT?"),
    ("E2", "format", "What must the student PICK?"),
    ("E3", "content_shape", "What NCERT MATERIAL is it built on?"),
    ("E4", "thinking", "What must the student do IN THEIR HEAD?"),
    ("E5", "trap", "Why does each WRONG OPTION attract?"),
]

_HEAD_FILL = PatternFill("solid", fgColor="1F3864")
_HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
_ALT_FILL = PatternFill("solid", fgColor="F4F6FA")
_GAP_FILL = PatternFill("solid", fgColor="FFF6E5")  # an empty why, so the hole is visible
_THIN = Side(style="thin", color="D5DBE5")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

COLS = [
    ("question_id", 15, False), ("subject", 11, False), ("year", 7, False),
    ("question_text", 68, True), ("options", 46, True), ("official_answer", 9, False),
    # Every question has a crop; B9 recovers a figure for those whose source dropped it. Without
    # these a row reads "[missing content]" while the picture sits on disk.
    *qa.ASSET_COLUMNS,
    ("E1 question type", 24, True), ("E1 why", 46, True),
    ("E2 format", 22, True), ("E2 why", 46, True),
    ("E3 content shape", 22, True), ("E3 why", 46, True),
    ("E4 thinking", 20, True), ("E4 why", 46, True),
    ("E5 trap", 26, True), ("E5 why", 46, True),
]
_WHY_COLS = {i for i, (h, _, _) in enumerate(COLS, 1) if h.endswith(" why")}


def load_axis_data(final_dir):
    """(tags, whys) -- question_id -> {axis: value}. Values joined with ';' when multi-valued."""
    tags = collections.defaultdict(dict)
    whys = collections.defaultdict(dict)
    for axis, tag, _ in AXES:
        path = os.path.join(final_dir, f"{tag}.json")
        with open(path, encoding="utf-8") as fh:
            doc = json.load(fh)
        for cat in doc.get("categories") or []:
            if not isinstance(cat, dict):
                continue
            name = cat.get("name")
            for qid in (cat.get("question_ids") or []):
                prev = tags[qid].get(axis)
                tags[qid][axis] = f"{prev};{name}" if prev else name
            for ex in (cat.get("examples") or []):
                if isinstance(ex, dict) and ex.get("question_id") and ex.get("why"):
                    whys[ex["question_id"]].setdefault(axis, ex["why"])
    return tags, whys


def load_records():
    out = {}
    d = rc.pp("papers", "records")
    for fn in sorted(os.listdir(d)):
        if not (fn.startswith("Q-") and fn.endswith(".json")):
            continue
        try:
            with open(os.path.join(d, fn), encoding="utf-8") as fh:
                rec = json.load(fh)
        except (OSError, json.JSONDecodeError):
            continue
        if isinstance(rec, dict) and rec.get("question_id"):
            out[rec["question_id"]] = rec
    return out


def question_text(rec):
    """Stem plus lists and statements -- the stem alone is boilerplate for match and A-R shapes."""
    parts = [(rec.get("stem_en") or "").strip()]
    for field, label in (("statements", ""), ("list_i", "List I"), ("list_ii", "List II")):
        entries = [e for e in (rec.get(field) or []) if isinstance(e, dict)]
        if entries:
            if label:
                parts.append(f"{label}:")
            parts.extend(f"({e.get('label')}) {e.get('text')}" for e in entries)
    return "\n".join(p for p in parts if p)


def options_text(rec):
    opts = {k: (v or "").strip() for k, v in sorted((rec.get("options_en") or {}).items())}
    return " | ".join(f"({k}) {v}" for k, v in opts.items() if v)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--final-dir", default=rc.pp("exports", "discovery", "final"))
    ap.add_argument("--out", default=rc.pp("exports", "discovery", "QUESTIONS.xlsx"))
    args = ap.parse_args()

    tags, whys = load_axis_data(args.final_dir)
    records = load_records()
    qids = sorted(tags)

    rows = []
    for qid in qids:
        rec = records.get(qid, {})
        row = [qid, rec.get("subject", ""), rec.get("year", ""),
               qa.question_text(rec) if rec else "",
               qa.options_text(rec, one_per_line=False) if rec else "",
               rec.get("official_answer", ""), *qa.asset_values(rec)]
        for axis, _, _ in AXES:
            row.append(tags[qid].get(axis, ""))
            row.append(whys.get(qid, {}).get(axis, ""))
        rows.append(row)

    wb = Workbook()
    wb.remove(wb.active)

    # ---- Read me
    ws = wb.create_sheet("Read me")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 106
    written = {a: sum(1 for q in qids if whys.get(q, {}).get(a)) for a, _, _ in AXES}
    tagged = {a: sum(1 for q in qids if tags[q].get(a)) for a, _, _ in AXES}
    no_record = [q for q in qids if q not in records]
    info = [
        ("QUESTIONS — the per-question sheet", ""),
        ("", ""),
        ("What this is", "One row per analysed question, carrying its value on each of the five "
                         "discovery axes plus the reason for that placement. CATEGORIES.xlsx lists "
                         "the categories; this lists the questions."),
        ("Questions", len(qids)),
        ("Figures", (lambda c: f"{c['declare_a_figure']} questions depend on a picture; "
                               f"{c['figure_recovered']} have one recovered, "
                               f"{c['figure_still_needed']} still needed. Questions whose text is "
                               f"complete carry no image path at all.")(
                   qa.coverage([records[q] for q in qids if q in records]))),
        ("", ""),
    ]
    for axis, tag, question in AXES:
        info.append((f"{axis} {tag}", f"{question}   —   {tagged[axis]} questions tagged, "
                                      f"{written[axis]} with a written why"))
    info += [
        ("", ""),
        ("Where values come from", "Inverted from exports/discovery/final/<tag>.json, where each "
                                   "category already carries a question_ids list. Read, not "
                                   "recomputed."),
        ("Multi-valued cells", "E1 and E5 can hold several values joined by ';'. E1 because some "
                               "questions are built two ways at once (a plain MCQ that is also a "
                               "figure to read); E5 because that tag runs once per WRONG OPTION, so "
                               "one question can carry several traps."),
        ("Empty why cells", f"Only {sum(written.values())} of {len(qids) * len(AXES)} why cells are "
                            "written (580 questions x 5 axes). The why text comes from each "
                            "category's `examples`, which were written to illustrate a category, not "
                            "to cover every question. Empty cells are tinted and left blank on "
                            "purpose: filling them with the category's definition would read as "
                            "per-question evidence while carrying none."),
        ("Figure file", "papers/figures/<QID>.png — written by step B9 when the question's own "
                        "source dropped its diagram. A tinted 'NEEDED' cell means the question "
                        "declares a figure and none has been recovered yet; a blank means the "
                        "question needs no figure."),
        ("Figure source", "Which legacy sheet row or Mathpix URL the recovered figure came from, "
                          "kept as provenance. The figure's own option numbering belongs to ITS "
                          "source paper, which may be a different booklet code from this question "
                          "— so never read the pictured (1)(2)(3)(4) as this question's options."),
        ("Question text", "From papers/records/. For match and assertion-reason shapes the stem "
                          "alone is boilerplate, so lists and statements are included — that is "
                          "where the discriminating content sits."),
        ("Missing records", f"{len(no_record)} tagged question(s) have no record in papers/records/"
                            + (": " + ", ".join(no_record[:10]) if no_record else "")),
    ]
    for label, value in info:
        ws.append([label, value])
    ws["A1"].font = Font(bold=True, size=15, color="1F3864")
    for r in range(3, len(info) + 2):
        ws.cell(row=r, column=1).font = Font(bold=True, size=10)
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top", wrap_text=True)

    # ---- Questions
    qs = wb.create_sheet("Questions")
    qs.append([c[0] for c in COLS])
    for i, (_, width, _) in enumerate(COLS, 1):
        qs.column_dimensions[get_column_letter(i)].width = width
    for cell in qs[1]:
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = _BORDER
    qs.row_dimensions[1].height = 30
    for n, row in enumerate(rows, start=2):
        qs.append(row)
        for i, (_, _, wrap) in enumerate(COLS, 1):
            cell = qs.cell(row=n, column=i)
            cell.alignment = Alignment(vertical="top", wrap_text=wrap)
            cell.border = _BORDER
            if COLS[i - 1][0] == "figure file" and str(cell.value or "").startswith("NEEDED"):
                cell.fill = _GAP_FILL
            elif i in _WHY_COLS and not cell.value:
                cell.fill = _GAP_FILL
            elif n % 2 == 0:
                cell.fill = _ALT_FILL
    qs.freeze_panes = "D2"
    qs.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(rows) + 1}"

    wb.save(args.out)
    print(f"wrote {os.path.relpath(args.out, rc.pp())}")
    print(f"  Questions: {len(rows)} rows x {len(COLS)} columns")
    for axis, tag, _ in AXES:
        print(f"  {axis} {tag:14s} tagged {tagged[axis]:4d}   why written {written[axis]:4d}")
    print(f"  why cells written {sum(written.values())} of {len(qids) * len(AXES)}")
    if no_record:
        print(f"  WARNING: {len(no_record)} tagged question(s) absent from papers/records/: "
              f"{no_record[:6]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
